import csv
import io
import json
from datetime import date, datetime, timedelta, timezone
from unittest.mock import Mock
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, func, text
from app.db.database import Database
from app.gsc.runtime import Runtime, get_runtime
from app.gsc.dates import bounds
from app.gsc.client import FixtureGSCClient, GoogleGSCClient
from app.gsc.repository import create_fixture_properties, properties_view
from app.keywords.core import change, week_dates, normalize
from app.keywords.models import RankObservation, KeywordMetric, KeywordSyncJob
from app.keywords.schemas import SyncInput
from app.keywords.gsc import fetch_exact
from app.main import app
from app.api.gsc_routes import CSRF_TOKEN


@pytest.fixture
def env(tmp_path):
    runtime = Runtime(Database(tmp_path / "kw.db"))
    app.dependency_overrides[get_runtime] = lambda: runtime
    client = TestClient(
        app, base_url="http://127.0.0.1:8000", headers={"X-CSRF-Token": CSRF_TOKEN}
    )
    yield client, runtime
    app.dependency_overrides.clear()
    runtime.keywords.close()
    runtime.sync.close()
    runtime.technical.close()
    runtime.db.engine.dispose()


def setup(env, site="xuong-may", start="2026-10-19"):
    c, r = env
    response = c.post(
        "/api/keyword-tracking/sets",
        json={"site_id": site, "name": "Bộ từ khóa tiếng Việt", "start_date": start},
    )
    assert response.status_code == 201, response.text
    ks = response.json()
    kw = c.post(
        "/api/keyword-tracking/keywords",
        json={
            "keyword_set_id": ks["id"],
            "keyword_text": "May balô giá rẻ",
            "target_url": "https://example.com/balo/",
        },
    ).json()
    return ks, kw


def rank(c, ks, kw, value, week=1, **extra):
    return c.post(
        "/api/keyword-tracking/ranks",
        json={
            "keyword_id": kw["id"],
            "week_number": week,
            "profile_id": ks["profiles"][0]["id"],
            "value": value,
            "timestamp_precision": "week_only",
            **extra,
        },
    )


def table(c, ks, **extra):
    params = {"set_id": ks["id"], "profile_id": ks["profiles"][0]["id"], **extra}
    response = c.get("/api/keyword-tracking/keywords", params=params)
    assert response.status_code == 200, response.text
    return response.json()


def test_management_duplicate_unicode_isolation_and_atomic_bulk(env):
    c, r = env
    ks, kw = setup(env)
    assert len(ks["weeks"]) == 35
    assert normalize("  MAy  balô giá RẺ  ") == normalize(kw["keyword_text"])
    assert (
        c.post(
            "/api/keyword-tracking/keywords",
            json={"keyword_set_id": ks["id"], "keyword_text": "  MAY balô giá rẻ  "},
        ).status_code
        == 409
    )
    second, _ = setup(env)
    group = c.post(
        "/api/keyword-tracking/groups",
        json={"keyword_set_id": ks["id"], "name": "Pillar", "color": "#895537"},
    ).json()
    assert (
        c.patch(
            "/api/keyword-tracking/keywords/" + kw["id"],
            json={"group_id": group["id"], "target_url": "https://example.com/new/"},
        ).status_code
        == 200
    )
    assert table(c, ks)["rows"][0]["group_id"] == group["id"]
    other = table(c, second)["rows"][0]
    response = c.post(
        "/api/keyword-tracking/keywords/bulk",
        json={"ids": [kw["id"], other["id"]], "changes": {"status": "archived"}},
    )
    assert response.status_code == 422
    assert table(c, ks)["total"] == 1
    assert (
        c.patch(
            "/api/keyword-tracking/keywords/" + kw["id"], json={"status": "archived"}
        ).status_code
        == 200
    )
    assert table(c, ks)["total"] == 0
    assert table(c, ks, status="archived")["total"] == 1


@pytest.mark.parametrize(
    "value,status,number",
    [
        ("1", "FOUND", 1),
        ("100", "FOUND", 100),
        ("NA", "NOT_FOUND", None),
        ("", "NOT_CHECKED", None),
    ],
)
def test_rank_states_persistence(value, status, number, env):
    c, r = env
    ks, kw = setup(env)
    res = rank(c, ks, kw, value)
    assert res.status_code == 201, res.text
    row = table(c, ks)["rows"][0]
    assert row["weeks"]["1"]["rank_status"] == status
    assert row["weeks"]["1"]["rank_value"] == number
    reopened = Database(r.db.path)
    reopened.migrate()
    with reopened.session() as s:
        assert s.get(RankObservation, res.json()["id"]).rank_status == status
    reopened.engine.dispose()


@pytest.mark.parametrize("value", ["0", "-1", "2.5", "abc", "101", "1e1"])
def test_reject_invalid_rank(value, env):
    c, r = env
    ks, kw = setup(env)
    assert rank(c, ks, kw, value).status_code == 422
    assert table(c, ks)["rows"][0]["weeks"] == {}


def test_latest_clear_history_bulk_rollback_and_calendar_lock(env):
    c, r = env
    ks, kw = setup(env)
    assert rank(c, ks, kw, "20").status_code == 201
    assert rank(c, ks, kw, "8").status_code == 201
    assert table(c, ks)["rows"][0]["weeks"]["1"]["count"] == 2
    assert rank(c, ks, kw, "").status_code == 201
    history = c.get(
        f"/api/keyword-tracking/keywords/{kw['id']}/history",
        params={"profile_id": ks["profiles"][0]["id"]},
    ).json()
    assert (
        len(history["rows"]) == 3 and history["rows"][0]["rank_status"] == "NOT_CHECKED"
    )
    response = c.post(
        "/api/keyword-tracking/ranks/bulk",
        json={
            "rows": [
                {
                    "keyword_id": kw["id"],
                    "profile_id": ks["profiles"][0]["id"],
                    "week_number": 2,
                    "value": "9",
                    "timestamp_precision": "week_only",
                },
                {
                    "keyword_id": kw["id"],
                    "profile_id": ks["profiles"][0]["id"],
                    "week_number": 3,
                    "value": "0",
                },
            ]
        },
    )
    assert response.status_code == 422
    assert "2" not in table(c, ks)["rows"][0]["weeks"]
    assert (
        c.patch(
            "/api/keyword-tracking/sets/" + ks["id"], json={"start_date": "2026-10-26"}
        ).status_code
        == 409
    )
    assert table(c, ks)["rows"][0]["weeks"]["1"]["timestamp_precision"] == "week_only"


@pytest.mark.parametrize(
    "a,b,state,delta",
    [
        ("20", "8", "improved", 12),
        ("8", "20", "dropped", -12),
        ("10", "10", "unchanged", 0),
        ("10", "NA", "lost", None),
        ("NA", "10", "entered", None),
        ("NA", "NA", "still_not_found", None),
        ("", "10", "new_observation", None),
        ("10", "", "awaiting", None),
    ],
)
def test_rank_changes(a, b, state, delta, env):
    c, r = env
    ks, kw = setup(env)
    rank(c, ks, kw, a)
    rank(c, ks, kw, b, 2)
    row = table(c, ks, week=2)["rows"][0]
    assert row["change"] == {"state": state, "value": delta}


def test_calendar_boundaries_profile_and_target_snapshots(env):
    c, r = env
    ks, kw = setup(env)
    assert week_dates("2026-10-26", 35) == ("2027-06-21", "2027-06-27")
    assert (
        rank(
            c,
            ks,
            kw,
            "9",
            timestamp_precision="exact",
            checked_at="2026-10-18T17:00:00Z",
        ).status_code
        == 201
    )
    assert (
        rank(
            c,
            ks,
            kw,
            "9",
            timestamp_precision="exact",
            checked_at="2026-10-18T16:59:59Z",
        ).status_code
        == 422
    )
    rank(c, ks, kw, "20")
    c.patch(
        "/api/keyword-tracking/keywords/" + kw["id"],
        json={"target_url": "https://example.com/new/"},
    )
    rank(c, ks, kw, "8", 2)
    row = table(c, ks, week=2)["rows"][0]
    assert row["change"]["state"] == "incomparable"
    assert row["weeks"]["1"]["target_url_snapshot"] == "https://example.com/balo/"
    assert (
        rank(c, ks, kw, "5", matched_url="https://elsewhere.example/").status_code
        == 422
    )


def test_pagination_filters_sql_sort_and_keyword_rename(env):
    c, r = env
    ks, kw = setup(env)
    rank(c, ks, kw, "8")
    assert table(c, ks, rank_filter="top10")["total"] == 1
    assert table(c, ks, rank_filter="unchecked")["total"] == 0
    assert table(c, ks, search="%")["total"] == 0
    assert (
        c.get(
            "/api/keyword-tracking/keywords",
            params={
                "set_id": ks["id"],
                "profile_id": ks["profiles"][0]["id"],
                "sort": "1; DROP TABLE keyword_sets",
            },
        ).status_code
        == 422
    )
    assert (
        c.patch(
            "/api/keyword-tracking/keywords/" + kw["id"],
            json={"keyword_text": "Từ khóa mới"},
        ).status_code
        == 200
    )
    assert (
        table(c, ks)["rows"][0]["weeks"]["1"]["profile_snapshot"]["keyword_text"]
        == "May balô giá rẻ"
    )


def test_gsc_exact_filter_aggregation_and_cache(env):
    c, r = env
    create_fixture_properties(r.db, FixtureGSCClient())
    prop = properties_view(r.db)[0]
    ks, kw = setup(env, prop.id)
    c.patch(
        "/api/keyword-tracking/keywords/" + kw["id"],
        json={
            "keyword_text": "túi xách thử nghiệm 1",
            "target_url": "https://fixture.example/trang-1",
        },
    )
    end = bounds()[1]
    data = SyncInput(
        keyword_set_id=ks["id"], date_start=end - timedelta(days=6), date_end=end
    )
    job = r.keywords.start(data, submit=False)
    assert r.keywords.start(data, submit=False)["id"] == job["id"]
    r.keywords.run(job["id"])
    row = table(
        c, ks, date_start=data.date_start.isoformat(), date_end=end.isoformat()
    )["rows"][0]
    assert row["metrics"]["status"] == "available"
    assert row["metrics"]["source"].startswith("fixture")
    assert row["metrics"]["average_position"] == 6.5
    page = SyncInput(**{**data.model_dump(), "scope": "page"})
    job = r.keywords.start(page, submit=False)
    r.keywords.run(job["id"])
    assert (
        table(
            c,
            ks,
            date_start=data.date_start.isoformat(),
            date_end=end.isoformat(),
            scope="page",
        )["rows"][0]["metrics"]["average_position"]
        == 7.5
    )
    with r.db.session() as s:
        assert s.scalar(select(func.count()).select_from(KeywordMetric)) == 2
    client = Mock()
    client.exact_query.return_value = {}
    fetch_exact(
        client,
        "sc-domain:test.com",
        "Từ Khóa",
        data.date_start.isoformat(),
        end.isoformat(),
        page="https://test.com/a",
    )
    body = client.exact_query.call_args.args[1]
    assert body["aggregationType"] == "byPage"
    assert body["dimensionFilterGroups"][0]["filters"] == [
        {"dimension": "query", "operator": "equals", "expression": "Từ Khóa"},
        {"dimension": "page", "operator": "equals", "expression": "https://test.com/a"},
    ]


def spreadsheet(env, format="csv", bad=False):
    from app.keywords.spreadsheets import HEADERS

    c, r = env
    ks, kw = setup(env)
    rows = [
        HEADERS,
        [
            "Từ khóa nhập",
            "Cate",
            "https://example.com/nhap/",
            "42",
            "100",
            "5.5",
            "NA",
            "8",
        ]
        + [""] * 33,
    ]
    if bad:
        rows.append(["Dòng lỗi", "", "", "", "", "", "0"] + [""] * 34)
    if format == "csv":
        buf = io.StringIO()
        csv.writer(buf).writerows(rows)
        content = buf.getvalue().encode("utf-8-sig")
    else:
        book = Workbook()
        for row in rows:
            book.active.append(row)
        out = io.BytesIO()
        book.save(out)
        book.close()
        content = out.getvalue()
    response = c.post(
        "/api/keyword-tracking/import",
        files={"file": ("test." + format, content)},
        data={
            "set_id": ks["id"],
            "profile_id": ks["profiles"][0]["id"],
            "mode": "full",
        },
    )
    assert response.status_code == 200, response.text
    return ks, response.json()


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_import_preview_commit_roundtrip_35_weeks_and_unicode(format, env):
    c, r = env
    ks, job = spreadsheet(env, format)
    assert job["data"]["total_rows"] == 1
    assert job["data"]["samples"][0]["week_1"] == "NA"
    assert table(c, ks)["total"] == 1  # Preview did not change keywords.
    response = c.post(
        "/api/keyword-tracking/import/" + job["id"] + "/commit", json={"confirm": True}
    )
    assert response.status_code == 200, response.text
    assert response.json()["imported"] == 1
    row = next(x for x in table(c, ks)["rows"] if x["keyword_text"] == "Từ khóa nhập")
    assert (
        row["weeks"]["1"]["rank_status"] == "NOT_FOUND"
        and row["weeks"]["2"]["rank_value"] == 8
    )
    assert row["weeks"]["1"]["checked_at"] is None and row["metrics"]["clicks"] is None
    response = c.get(
        "/api/keyword-tracking/export",
        params={
            "set_id": ks["id"],
            "profile_id": ks["profiles"][0]["id"],
            "format": format,
        },
    )
    assert response.status_code == 200, response.text
    if format == "csv":
        assert response.content.startswith(b"\xef\xbb\xbf")
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
        assert len(rows[0]) == 41
    else:
        book = load_workbook(io.BytesIO(response.content))
        assert book.active.max_column == 41
        assert "Metadata" in book.sheetnames
        book.close()
    assert (
        c.post(
            "/api/keyword-tracking/import/" + job["id"] + "/commit",
            json={"confirm": True},
        ).status_code
        == 409
    )


def test_import_errors_explicit_acceptance_formula_and_security(env):
    c, r = env
    ks, job = spreadsheet(env, bad=True)
    assert job["data"]["errors"]
    assert (
        c.post(
            "/api/keyword-tracking/import/" + job["id"] + "/commit",
            json={"confirm": True},
        ).status_code
        == 422
    )
    assert (
        c.post(
            "/api/keyword-tracking/import/" + job["id"] + "/commit",
            json={"confirm": True, "accept_rejected": True},
        ).json()["skipped"]
        == 1
    )
    from app.keywords.spreadsheets import sanitize

    assert sanitize('=HYPERLINK("x")').startswith("'")
    assert sanitize(" @SUM(1)").startswith("'")
    assert (
        c.post(
            "/api/keyword-tracking/keywords",
            json={
                "keyword_set_id": ks["id"],
                "keyword_text": "x",
                "target_url": "javascript:alert(1)",
            },
        ).status_code
        == 422
    )
    assert (
        c.post(
            "/api/keyword-tracking/ranks", json={}, headers={"X-CSRF-Token": "bad"}
        ).status_code
        == 403
    )
    assert (
        c.post(
            "/api/keyword-tracking/sets",
            json={},
            headers={"Origin": "https://evil.example"},
        ).status_code
        == 403
    )


def test_move_with_history_preserves_snapshots_and_rejects_different_calendar(env):
    c, r = env
    ks, kw = setup(env)
    rank(c, ks, kw, "8")
    dest, _ = setup(env)
    response = c.post(
        "/api/keyword-tracking/keywords/" + kw["id"] + "/move",
        json={"keyword_set_id": dest["id"], "confirm": True},
    )
    assert (
        response.status_code == 409
    )  # Destination already tracks that normalized keyword.
    other = c.post(
        "/api/keyword-tracking/sets",
        json={"site_id": "xuong-may", "name": "Bộ mới", "start_date": ks["start_date"]},
    ).json()
    response = c.post(
        "/api/keyword-tracking/keywords/" + kw["id"] + "/move",
        json={"keyword_set_id": other["id"], "confirm": True},
    )
    assert response.status_code == 200, response.text
    new = c.get("/api/keyword-tracking/sets/" + other["id"]).json()
    row = table(c, new)["rows"][0]
    assert row["id"] == kw["id"]
    assert row["weeks"]["1"]["rank_value"] == 8
    assert row["weeks"]["1"]["target_url_snapshot"] == kw["target_url"]
    different = c.post(
        "/api/keyword-tracking/sets",
        json={"site_id": "xuong-may", "name": "Khác lịch", "start_date": "2026-10-26"},
    ).json()
    assert (
        c.post(
            "/api/keyword-tracking/keywords/" + kw["id"] + "/move",
            json={"keyword_set_id": different["id"], "confirm": True},
        ).status_code
        == 422
    )


def test_delete_confirmations_cascade_and_foreign_keys(env):
    c, r = env
    ks, kw = setup(env)
    rank(c, ks, kw, "8")
    assert (
        c.post(
            "/api/keyword-tracking/sets/" + ks["id"] + "/delete",
            json={"confirm": False},
        ).status_code
        == 422
    )
    assert (
        c.post(
            "/api/keyword-tracking/sets/" + ks["id"] + "/delete", json={"confirm": True}
        ).status_code
        == 200
    )
    with r.db.session() as s:
        assert s.scalar(select(func.count()).select_from(RankObservation)) == 0
        assert s.execute(text("PRAGMA foreign_key_check")).all() == []


def test_gsc_missing_cache_identity_retry_no_duplicate_calls_and_mismatch(env):
    c, r = env
    create_fixture_properties(r.db, FixtureGSCClient())
    p = properties_view(r.db)[0]
    ks, kw = setup(env, p.id)
    c.patch(
        "/api/keyword-tracking/keywords/" + kw["id"],
        json={
            "keyword_text": "túi xách thử nghiệm 1",
            "target_url": "https://fixture.example/other",
        },
    )
    end = bounds()[1]
    args = dict(
        keyword_set_id=ks["id"], date_start=end - timedelta(days=6), date_end=end
    )
    client = FixtureGSCClient()
    spy = Mock(wraps=client)
    r.keywords.client_factory = lambda prop: spy
    job = r.keywords.start(SyncInput(**args), submit=False)
    r.keywords.run(job["id"])
    assert spy.exact_query.call_count == 2
    row = table(c, ks, date_start=args["date_start"], date_end=end)["rows"][0]
    assert "url_mismatch" in row["indicators"]
    assert row["target_url"] == "https://fixture.example/other"
    second = r.keywords.start(SyncInput(**args), submit=False)
    r.keywords.run(second["id"])
    assert spy.exact_query.call_count == 2
    c.patch(
        "/api/keyword-tracking/keywords/" + kw["id"],
        json={"keyword_text": "query không có dữ liệu"},
    )
    assert (
        table(c, ks, date_start=args["date_start"], date_end=end)["rows"][0]["metrics"][
            "status"
        ]
        == "not_synced"
    )
    job = r.keywords.start(SyncInput(**args), submit=False)
    r.keywords.run(job["id"])
    row = table(c, ks, date_start=args["date_start"], date_end=end)["rows"][0]
    assert (
        row["metrics"]["status"] == "no_rows"
        and row["metrics"]["average_position"] is None
        and row["metrics"]["clicks"] is None
    )
    assert (
        c.post(
            "/api/keyword-tracking/gsc-sync",
            json={**{k: str(v) for k, v in args.items()}, "date_start": "2020-01-01"},
        ).status_code
        == 422
    )


def test_existing_daily_query_cache_reuse_and_insufficient_coverage(env):
    from app.db.models import Checkpoint, MetricFact
    from app.gsc.dates import now
    from app.keywords.gsc import cached_daily

    c, r = env
    create_fixture_properties(r.db, FixtureGSCClient())
    p = properties_view(r.db)[0]
    end = bounds()[1]
    start = end - timedelta(days=1)
    with r.db.session.begin() as s:
        for i, day in enumerate([start, end]):
            s.add(
                Checkpoint(
                    property_id=p.id,
                    date=day.isoformat(),
                    search_type="web",
                    dataset_type="query",
                    status="complete",
                    completeness="api_response",
                    fetched_rows=1,
                    aggregation_type="byProperty",
                    last_successful_sync_at=now(),
                )
            )
            s.add(
                MetricFact(
                    property_id=p.id,
                    date=day.isoformat(),
                    search_type="web",
                    dataset_type="query",
                    dimension_value="Từ Khóa",
                    aggregation_type="byProperty",
                    clicks=1,
                    impressions=10 if not i else 30,
                    ctr=0.1,
                    position=10 if not i else 20,
                    data_state="final",
                    fetched_at=now(),
                )
            )
    with r.db.session() as s:
        m = cached_daily(s, p, "Từ Khóa", start.isoformat(), end.isoformat())
        assert m["average_position"] == 17.5
        assert m["ctr"] == 0.05
        assert (
            cached_daily(
                s,
                p,
                "Từ Khóa",
                (start - timedelta(days=1)).isoformat(),
                end.isoformat(),
            )
            is None
        )
        assert (
            cached_daily(s, p, "từ khóa", start.isoformat(), end.isoformat()) is None
        )  # Never normalize Google's exact query expression.


def test_gsc_api_failure_keeps_cached_metrics_and_recovery(env):
    from app.gsc.errors import GSCError

    c, r = env
    create_fixture_properties(r.db, FixtureGSCClient())
    p = properties_view(r.db)[0]
    ks, kw = setup(env, p.id)
    c.patch(
        "/api/keyword-tracking/keywords/" + kw["id"],
        json={"keyword_text": "túi xách thử nghiệm 1"},
    )
    end = bounds()[1]
    data = SyncInput(
        keyword_set_id=ks["id"], date_start=end - timedelta(days=6), date_end=end
    )
    job = r.keywords.start(data, submit=False)
    r.keywords.run(job["id"])
    bad = Mock()
    bad.exact_query.side_effect = GSCError(
        "permission_denied", "Quyền property đã bị thu hồi.", 403
    )
    r.keywords.client_factory = lambda p: bad
    job = r.keywords.start(
        SyncInput(**{**data.model_dump(), "refresh": True}), submit=False
    )
    r.keywords.run(job["id"])
    with r.db.session() as s:
        assert s.get(KeywordSyncJob, job["id"]).status == "failed"
    assert (
        table(c, ks, date_start=data.date_start, date_end=end)["rows"][0]["metrics"][
            "status"
        ]
        == "available"
    )
    job = r.keywords.start(data, submit=False)
    r.keywords.recover()
    with r.db.session() as s:
        assert s.get(KeywordSyncJob, job["id"]).status == "interrupted"


def test_custom_mapping_duplicates_separate_campaign_and_archive_history(env):
    c, r = env
    ks, kw = setup(env)

    def preview(content, mapping="{}", duplicates="skip"):
        response = c.post(
            "/api/keyword-tracking/import",
            files={"file": ("a.csv", content.encode())},
            data={
                "set_id": ks["id"],
                "profile_id": ks["profiles"][0]["id"],
                "mapping": mapping,
                "duplicates": duplicates,
            },
        )
        assert response.status_code == 200, response.text
        return response.json()

    first = preview("Custom,Result\nTừ khóa mới,NA\n")
    assert first["data"]["errors"]
    job = preview(
        "Custom,Result\nTừ khóa mới,NA\n",
        json.dumps({"keyword_text": "Custom", "week_1": "Result"}),
    )
    assert (
        c.post(
            "/api/keyword-tracking/import/" + job["id"] + "/commit",
            json={"confirm": True},
        ).json()["imported"]
        == 1
    )
    job = preview("Từ khóa,Tuần 1\nMay balô giá rẻ,10\n", duplicates="separate")
    response = c.post(
        "/api/keyword-tracking/import/" + job["id"] + "/commit", json={"confirm": True}
    ).json()
    assert response["keyword_set_id"] != ks["id"]
    rank(c, ks, kw, "5")
    c.patch("/api/keyword-tracking/keywords/" + kw["id"], json={"status": "archived"})
    history = c.get(
        "/api/keyword-tracking/keywords/" + kw["id"] + "/history",
        params={"profile_id": ks["profiles"][0]["id"]},
    ).json()
    assert history["total"] == 1


def test_xlsx_formula_zip_and_upload_limits(env):
    c, r = env
    ks, kw = setup(env)
    b = Workbook()
    b.active.append(["Từ khóa", "Tuần 1"])
    b.active.append(["=1+1", 5])
    buf = io.BytesIO()
    b.save(buf)
    b.close()
    args = {"set_id": ks["id"], "profile_id": ks["profiles"][0]["id"]}
    assert (
        c.post(
            "/api/keyword-tracking/import",
            files={"file": ("formula.xlsx", buf.getvalue())},
            data=args,
        ).status_code
        == 422
    )
    assert (
        c.post(
            "/api/keyword-tracking/import",
            files={"file": ("large.csv", b"a" * 102401)},
            data=args,
        ).status_code
        == 422
    )
    assert (
        c.post(
            "/api/keyword-tracking/import",
            files={"file": ("bad.xlsx", b"not a zip")},
            data=args,
        ).status_code
        == 422
    )
    assert (
        c.post(
            "/api/keyword-tracking/keywords",
            json={
                "keyword_set_id": ks["id"],
                "keyword_text": "unsafe",
                "target_url": "https://user:pass@example.com/",
            },
        ).status_code
        == 422
    )


def test_upgrade_from_v03_preserves_gsc_fact_values_and_adds_only_tables(tmp_path):
    from alembic import command
    from alembic.config import Config
    from app.db.database import BACKEND

    db = Database(tmp_path / "upgrade.db")
    config = Config()
    config.set_main_option("script_location", str(BACKEND / "migrations"))
    with db.engine.begin() as conn:
        config.attributes["connection"] = conn
        command.upgrade(config, "002_technical_audits")
        conn.exec_driver_sql(
            "INSERT INTO google_connections VALUES ('old','Existing user','fixture','fixture','none','stamp','stamp')"
        )
        conn.exec_driver_sql(
            "INSERT INTO gsc_properties VALUES ('old-property','old','sc-domain:existing.example','Existing','domain','siteOwner','fixture',1,'stamp','stamp')"
        )
        conn.exec_driver_sql(
            "INSERT INTO gsc_daily_metrics (property_id,date,search_type,dataset_type,dimension_value,aggregation_type,clicks,impressions,ctr,position,data_state,fetched_at) VALUES ('old-property','2026-10-10','web','site','','byProperty',17,100,.17,9.5,'final','stamp')"
        )
        before = conn.exec_driver_sql("SELECT * FROM gsc_daily_metrics").all()
    db.migrate()
    db.migrate()
    with db.engine.connect() as conn:
        assert conn.exec_driver_sql("SELECT * FROM gsc_daily_metrics").all() == before
        assert (
            conn.exec_driver_sql("SELECT version_num FROM alembic_version").scalar()
            == "005_keyword_serp_research"
        )
        assert conn.exec_driver_sql("PRAGMA integrity_check").scalar() == "ok"
        assert conn.exec_driver_sql("PRAGMA foreign_key_check").all() == []
    db.engine.dispose()


def test_export_formula_is_text_and_na_and_blank_remain_distinct(env):
    c, r = env
    ks, kw = setup(env)
    c.patch(
        "/api/keyword-tracking/keywords/" + kw["id"],
        json={"keyword_text": '=HYPERLINK("https://example.com")'},
    )
    rank(c, ks, kw, "NA", 1)
    rank(c, ks, kw, "1", 2)
    rank(c, ks, kw, "", 3)
    args = {"set_id": ks["id"], "profile_id": ks["profiles"][0]["id"]}
    csvdata = c.get(
        "/api/keyword-tracking/export", params={**args, "format": "csv"}
    ).content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(csvdata)))
    assert rows[1][0].startswith("'=") and rows[1][6:9] == ["NA", "1", ""]
    response = c.get("/api/keyword-tracking/export", params={**args, "format": "xlsx"})
    book = load_workbook(io.BytesIO(response.content))
    assert book.active["A2"].data_type == "s" and book.active["A2"].value.startswith(
        "'="
    )
    assert (
        book.active["G2"].value == "NA"
        and book.active["H2"].value == 1
        and book.active["I2"].value is None
    )
    book.close()


def test_profile_edits_and_missing_previous_week_never_invent_delta(env):
    c, r = env
    ks, kw = setup(env)
    rank(c, ks, kw, "20", 1)
    rank(c, ks, kw, "8", 3)
    assert table(c, ks, week=3)["rows"][0]["change"]["state"] == "new_observation"
    p = ks["profiles"][0]
    response = c.patch(
        "/api/keyword-tracking/profiles/" + p["id"],
        json={
            "keyword_set_id": ks["id"],
            "name": p["name"],
            **p["config"],
            "device": "mobile",
        },
    )
    assert response.status_code == 200, response.text
    rank(c, ks, kw, "5", 4)
    row = table(c, ks, week=4)["rows"][0]
    assert row["change"]["state"] == "incomparable"
    assert row["weeks"]["3"]["profile_snapshot"]["device"] == "desktop"
    assert row["weeks"]["4"]["profile_snapshot"]["device"] == "mobile"


def test_move_does_not_merge_distinct_profiles_with_identical_configuration(env):
    c, r = env
    ks, kw = setup(env)
    rank(c, ks, kw, "8")
    p = ks["profiles"][0]
    second = c.post(
        "/api/keyword-tracking/profiles",
        json={"keyword_set_id": ks["id"], "name": p["name"], **p["config"]},
    ).json()
    assert rank(c, ks, kw, "20", profile_id=second["id"]).status_code == 201
    dest = c.post(
        "/api/keyword-tracking/sets",
        json={
            "site_id": "xuong-may",
            "name": "Bộ nhận",
            "start_date": ks["start_date"],
        },
    ).json()
    response = c.post(
        "/api/keyword-tracking/keywords/" + kw["id"] + "/move",
        json={"keyword_set_id": dest["id"], "confirm": True},
    )
    assert response.status_code == 200, response.text
    detail = c.get("/api/keyword-tracking/sets/" + dest["id"]).json()
    assert len(detail["profiles"]) == 2
    results = [
        table(c, detail, profile_id=profile["id"])["rows"][0]["weeks"]["1"][
            "rank_value"
        ]
        for profile in detail["profiles"]
    ]
    assert set(results) == {8, 20}


def test_import_duplicate_update_appends_history_and_keeps_original_target_snapshot(
    env,
):
    c, r = env
    ks, kw = setup(env)
    rank(c, ks, kw, "8")
    content = "Từ khóa,Nhóm,URL,Tuần 1,Tuần 2\nMay balô giá rẻ,Nhóm mới,https://example.com/new/,20,NA\n"
    preview = c.post(
        "/api/keyword-tracking/import",
        files={"file": ("update.csv", content.encode())},
        data={
            "set_id": ks["id"],
            "profile_id": ks["profiles"][0]["id"],
            "duplicates": "update",
            "mode": "history",
        },
    ).json()
    assert preview["data"]["warnings"]
    response = c.post(
        "/api/keyword-tracking/import/" + preview["id"] + "/commit",
        json={"confirm": True},
    )
    assert response.status_code == 200, response.text
    row = table(c, ks, week=2)["rows"][0]
    assert row["id"] == kw["id"]
    assert row["target_url"] == "https://example.com/new/"
    assert (
        row["weeks"]["1"]["rank_value"] == 20
        and row["weeks"]["1"]["count"] == 2
        and row["weeks"]["2"]["rank_status"] == "NOT_FOUND"
    )
    history = c.get(
        "/api/keyword-tracking/keywords/" + kw["id"] + "/history",
        params={"profile_id": ks["profiles"][0]["id"]},
    ).json()
    assert history["total"] == 3
    assert (
        history["rows"][-1]["target_url_snapshot"] == "https://example.com/balo/"
        and history["rows"][-1]["rank_value"] == 8
    )
