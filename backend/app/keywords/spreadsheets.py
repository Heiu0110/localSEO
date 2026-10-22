import csv
import io
import math
import zipfile
from pathlib import PurePath
from openpyxl import load_workbook, Workbook
from sqlalchemy import select
from .core import *
from .schemas import KeywordInput, RankInput

HEADERS = ["Từ khóa", "Nhóm", "URL", "Click", "Impression", "Vị trí trung bình"] + [
    f"Tuần {n}" for n in range(1, 36)
]
FIELDS = [
    "keyword_text",
    "group",
    "target_url",
    "clicks",
    "impressions",
    "position",
] + [f"week_{n}" for n in range(1, 36)]


def read_file(content, filename):
    if len(content) > 100 * 1024:
        fail("Tệp nhập tối đa 100 KiB. Chia tệp lớn thành nhiều phần.")
    suffix = PurePath(filename).suffix.lower()
    try:
        if suffix == ".csv":
            text = content.decode("utf-8-sig")
            sample = text[:4096]
            delimiter = (
                ";"
                if sample.splitlines()[0].count(";") > sample.splitlines()[0].count(",")
                else ","
            )
            rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
        elif suffix == ".xlsx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                infos = archive.infolist()
                if (
                    len(infos) > 100
                    or sum(x.file_size for x in infos) > 8 * 1024 * 1024
                    or any(
                        x.flag_bits & 1 or ".." in x.filename.split("/") for x in infos
                    )
                ):
                    fail("XLSX vượt giới hạn giải nén hoặc không hợp lệ.")
                if any(
                    "vbaProject" in x.filename or "externalLinks/" in x.filename
                    for x in infos
                ):
                    fail("Không hỗ trợ workbook có macro hoặc liên kết ngoài.")
            book = load_workbook(
                io.BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
            try:
                sheet = book.worksheets[0]
                rows = []
                if (sheet.max_row or 0) > 1001 or (sheet.max_column or 0) > 100:
                    fail("XLSX vượt 1.000 dòng dữ liệu hoặc 100 cột.")
                for row in sheet.iter_rows(
                    max_row=min(sheet.max_row or 1002, 1002),
                    max_col=min(sheet.max_column or 100, 100),
                ):
                    if len(rows) > 1000:
                        fail("Mỗi tệp tối đa 1.000 dòng dữ liệu.")
                    if len(row) > 100:
                        fail("Mỗi tệp tối đa 100 cột.")
                    if any(c.data_type == "f" for c in row):
                        fail(
                            "Không chấp nhận công thức trong tệp nhập. Dán giá trị rồi thử lại."
                        )
                    rows.append([c.value for c in row])
            finally:
                book.close()
        else:
            fail("Chỉ hỗ trợ CSV UTF-8 và XLSX.")
    except GSCError:
        raise
    except Exception:
        fail("Không đọc được tệp. Kiểm tra định dạng CSV UTF-8 hoặc XLSX không mã hóa.")
    if not rows or len(rows) > 1001 or any(len(row) > 100 for row in rows):
        fail("Tệp cần header và tối đa 1.000 dòng, 100 cột.")
    headers = [str(x or "").strip() for x in rows[0]]
    if len(set(headers)) != len(headers) or not all(headers):
        fail("Header cần có tên và không trùng nhau.")
    result = []
    for row in rows[1:]:
        if any(x is not None and str(x).strip() for x in row):
            values = [str(x).strip() if x is not None else "" for x in row]
            if any(len(x) > 4000 for x in values):
                fail("Ô dữ liệu vượt 4.000 ký tự.")
            result.append(dict(zip(headers, values)))
    return headers, result


def preview(session, ks, profile, content, filename, mode, duplicates, mapping):
    if mode not in ("keywords", "history", "full") or duplicates not in (
        "skip",
        "update",
        "separate",
    ):
        fail("Chế độ nhập không hợp lệ.")
    headers, raw = read_file(content, filename)
    detected = {
        field: next(
            (
                h
                for h in headers
                if normalize(h) in (normalize(label), normalize(field))
            ),
            None,
        )
        for field, label in zip(FIELDS, HEADERS)
    }
    mapping = {**detected, **mapping}
    if any(k not in FIELDS or (v and v not in headers) for k, v in mapping.items()):
        fail("Ánh xạ cột không hợp lệ.")
    existing = {
        x.normalized_keyword
        for x in session.scalars(
            select(TrackedKeyword).where(TrackedKeyword.keyword_set_id == ks.id)
        )
    }
    seen = set()
    rows = []
    errors = []
    warnings = []
    for index, source in enumerate(raw, 2):
        row = {k: source.get(v, "") if v else "" for k, v in mapping.items()}
        row["row"] = index
        issues = []
        norm = normalize(row["keyword_text"])
        if not norm or len(row["keyword_text"]) > 500:
            issues.append("Từ khóa trống hoặc quá 500 ký tự.")
        if len(row["group"]) > 100:
            issues.append("Tên nhóm quá 100 ký tự.")
        try:
            row["target_url"] = safe_url(row["target_url"])
        except GSCError as e:
            issues.append(e.message)
        if norm in seen:
            issues.append("Từ khóa trùng trong tệp; tách campaign trước khi nhập.")
        seen.add(norm)
        row["duplicate"] = norm in existing
        if row["duplicate"]:
            warnings.append({"row": index, "message": "Trùng từ khóa: " + duplicates})
        if mode != "keywords":
            for n in range(1, 36):
                try:
                    parse_rank(row[f"week_{n}"], profile.config["search_depth"])
                except GSCError:
                    issues.append(f"Tuần {n}: rank không hợp lệ.")
        if mode == "full":
            for key in ("clicks", "impressions", "position"):
                try:
                    if row[key] and (
                        not math.isfinite(float(row[key]))
                        or float(row[key]) < 0
                        or (key != "position" and not float(row[key]).is_integer())
                    ):
                        raise ValueError()
                except ValueError:
                    issues.append(f"{key}: số không hợp lệ.")
        row["valid"] = not issues
        rows.append(row)
        errors.extend({"row": index, "message": message} for message in issues)
    data = {
        "headers": headers,
        "mapping": mapping,
        "rows": rows,
        "errors": errors,
        "warnings": warnings,
        "mode": mode,
        "duplicates": duplicates,
        "profile_id": profile.id,
        "calendar": ks.start_date,
        "profile_config": profile.config,
        "raw_rows": raw,
        "metrics_source": "historical_spreadsheet_unverified",
    }
    job = KeywordImport(
        id=uid(),
        keyword_set_id=ks.id,
        filename=filename.replace("\\", "/").split("/")[-1][:200],
        status="preview",
        data=data,
        created_at=now(),
    )
    session.add(job)
    session.flush()
    return job


def commit(session, job, confirm, accept_rejected=False):
    if not confirm:
        fail("Cần xác nhận bản xem trước trước khi nhập.")
    if job.status != "preview":
        fail("Bản nhập này đã được xử lý.", "import_already_committed", 409)
    data = job.data
    ks = require(session, KeywordSet, job.keyword_set_id)
    profile = require(session, SearchProfile, data["profile_id"])
    if ks.start_date != data["calendar"] or profile.config != data["profile_config"]:
        fail("Lịch hoặc profile đã đổi. Tạo bản xem trước mới.")
    if not data["mapping"].get("keyword_text") or not any(
        row["valid"] for row in data["rows"]
    ):
        fail("Chọn ánh xạ Từ khóa và cần ít nhất một dòng hợp lệ trước khi nhập.")
    if data["errors"] and not accept_rejected:
        fail(
            "Có dòng bị từ chối. Xác nhận bỏ qua các dòng lỗi hoặc sửa tệp rồi xem trước lại."
        )
    dest = ks
    if data["duplicates"] == "separate":
        dest = KeywordSet(
            id=uid(),
            site_id=ks.site_id,
            property_id=ks.property_id,
            name=(ks.name + " · campaign nhập")[:200],
            description="Bộ riêng từ tệp " + job.filename,
            start_date=ks.start_date,
            week_start_day=ks.week_start_day,
            timezone=ks.timezone,
            status="active",
            settings=ks.settings.copy(),
            created_at=now(),
            updated_at=now(),
        )
        session.add(dest)
        session.flush()
        ensure_weeks(session, dest)
        profile = SearchProfile(
            id=uid(),
            keyword_set_id=dest.id,
            name=profile.name,
            config=profile.config.copy(),
            created_at=now(),
            updated_at=now(),
        )
        session.add(profile)
        session.flush()
    groups = {
        g.name: g
        for g in session.scalars(
            select(KeywordGroup).where(KeywordGroup.keyword_set_id == dest.id)
        )
    }
    imported = skipped = 0
    for row in data["rows"]:
        if not row["valid"]:
            skipped += 1
            continue
        old = session.scalar(
            select(TrackedKeyword).where(
                TrackedKeyword.keyword_set_id == dest.id,
                TrackedKeyword.normalized_keyword == normalize(row["keyword_text"]),
            )
        )
        if old and data["duplicates"] == "skip":
            skipped += 1
            continue
        group = None
        if row["group"]:
            if row["group"] not in groups:
                group = KeywordGroup(
                    id=uid(),
                    keyword_set_id=dest.id,
                    name=row["group"],
                    color="#946447",
                    sort_order=0,
                )
                session.add(group)
                session.flush()
                groups[group.name] = group
            group = groups[row["group"]]
        if old:
            from .schemas import KeywordPatch

            kw = patch_keyword(
                session,
                old,
                KeywordPatch(
                    group_id=group.id if group else None,
                    target_url=row["target_url"],
                    match_mode="exact_url" if row["target_url"] else "domain",
                ),
            )
        else:
            kw = create_keyword(
                session,
                KeywordInput(
                    keyword_set_id=dest.id,
                    keyword_text=row["keyword_text"],
                    group_id=group.id if group else None,
                    target_url=row["target_url"],
                ),
            )
        if data["mode"] != "keywords":
            for n in range(1, 36):
                if row[f"week_{n}"]:
                    add_rank(
                        session,
                        RankInput(
                            keyword_id=kw.id,
                            week_number=n,
                            profile_id=profile.id,
                            value=row[f"week_{n}"],
                            notes="Nhập lịch sử: chỉ biết tuần; " + job.filename,
                        ),
                        imported=True,
                    )
        imported += 1
    job.status = "imported"
    job.data = {
        **data,
        "imported_rows": imported,
        "rejected_rows": skipped,
        "destination_set_id": dest.id,
        "committed_at": now(),
    }
    activity(
        session,
        dest.id,
        "import_committed",
        {"import_id": job.id, "imported": imported, "skipped": skipped},
    )
    return {"imported": imported, "skipped": skipped, "keyword_set_id": dest.id}


def sanitize(value):
    if isinstance(value, str) and value.lstrip().startswith(
        ("=", "+", "-", "@", "\t", "\r")
    ):
        return "'" + value
    return value


def export_file(rows, metadata, format):
    table = [HEADERS]
    for row in rows:
        m = row["metrics"]
        values = [
            row["keyword_text"],
            row.get("group_name", ""),
            row["target_url"],
            m.get("clicks"),
            m.get("impressions"),
            m.get("average_position"),
        ]
        for n in range(1, 36):
            obs = row["weeks"].get(str(n))
            values.append(
                obs["rank_value"]
                if obs and obs["rank_status"] == "FOUND"
                else "NA" if obs and obs["rank_status"] == "NOT_FOUND" else ""
            )
        table.append([sanitize(x) for x in values])
    if format == "csv":
        out = io.StringIO(newline="")
        csv.writer(out).writerows(table)
        return ("\ufeff" + out.getvalue()).encode("utf-8"), "text/csv; charset=utf-8"
    book = Workbook()
    sheet = book.active
    sheet.title = "Theo dõi từ khóa"
    for row in table:
        sheet.append(row)
    sheet.freeze_panes = "G2"
    sheet.auto_filter.ref = sheet.dimensions
    for col in ("A", "B", "C"):
        sheet.column_dimensions[col].width = {"A": 35, "B": 18, "C": 45}[col]
    info = book.create_sheet("Metadata")
    for k, v in metadata.items():
        info.append([sanitize(k), sanitize(str(v))])
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    return (
        buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
